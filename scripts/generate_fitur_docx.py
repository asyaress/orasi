# -*- coding: utf-8 -*-
"""Generate dokumentasi fitur Portal Orasi UNMUL — desain modern tema Orasi."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"
LOGO_DIR = PUBLIC / "logo"
OUTPUT = ROOT / "docs" / "Dokumentasi-Fitur-Portal-Orasi-UNMUL.docx"

# Tema Orasi (sama dengan portal)
C = {
    "navy": "18213A",
    "navy_dark": "0F1322",
    "gold": "EFB12C",
    "gold_soft": "F6C85C",
    "gold_pale": "FFF6D6",
    "blue": "2F8CF6",
    "muted": "667085",
    "light": "F4F6FA",
    "white": "FFFFFF",
    "border": "E2E8F0",
}

BRAND_LOGOS = [
    ("tut-wuri-20260408145730-756785.png", "Tut Wuri Handayani"),
    ("unmul-20260408145731-e033c2.png", "UNMUL"),
    ("blu-20260408145731-85748a.png", "BLU"),
    ("dies-natalis-20260408145732-edfb3f.png", "Dies Natalis"),
    ("diktisaintek-20260408145732-367e09.png", "Diktisaintek"),
    ("logo-unggul-20260408145732-41a84d.png", "Unggul UNMUL"),
]

FONT = "Segoe UI"


def rgb(hex_code: str) -> RGBColor:
    h = hex_code.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_cell_shading(cell, fill_hex: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def set_cell_borders(cell, color: str = C["border"], size: str = "6") -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tc_pr.append(borders)


def set_table_width(table, width_inches: float) -> None:
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(width_inches)


def remove_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def style_run(run, *, size=11, bold=False, color=C["navy"], italic=False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.insert(0, r_fonts)


def add_paragraph_text(
    doc: Document,
    text: str,
    *,
    size=11,
    bold=False,
    color=C["navy"],
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_after=6,
    italic=False,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color, italic=italic)


def add_accent_line(doc: Document, height_pt: float = 4) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, 6.5)
    remove_table_borders(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, C["gold"])
    set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
    cell.paragraphs[0].paragraph_format.space_before = Pt(0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(height_pt)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_themed_heading(doc: Document, text: str, level: int = 1) -> None:
    if level == 1:
        add_accent_line(doc)
        table = doc.add_table(rows=1, cols=1)
        set_table_width(table, 6.5)
        remove_table_borders(table)
        cell = table.rows[0].cells[0]
        set_cell_shading(cell, C["navy"])
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        run = p.add_run(text.upper())
        style_run(run, size=14, bold=True, color=C["white"])
        doc.add_paragraph().paragraph_format.space_after = Pt(8)
        return

    if level == 2:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        style_run(run, size=12, bold=True, color=C["navy"])
        # gold underline bar
        bar = doc.add_table(rows=1, cols=1)
        set_table_width(bar, 1.6)
        remove_table_borders(bar)
        bc = bar.rows[0].cells[0]
        set_cell_shading(bc, C["gold"])
        bc.paragraphs[0].paragraph_format.space_after = Pt(2)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    style_run(run, size=11, bold=True, color=C["blue"])


def add_info_box(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, 6.5)
    remove_table_borders(table)
    accent, body = table.rows[0].cells[0], table.rows[0].cells[1]
    accent.width = Inches(0.08)
    body.width = Inches(6.42)
    set_cell_shading(accent, C["gold"])
    set_cell_shading(body, C["gold_pale"])
    set_cell_margins(body, top=100, bottom=100, left=160, right=160)
    set_cell_borders(body, C["gold_soft"], "4")
    run = body.paragraphs[0].add_run(text)
    style_run(run, size=10, color=C["navy"])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        table = doc.add_table(rows=1, cols=2)
        set_table_width(table, 6.5)
        remove_table_borders(table)
        dot, text_cell = table.rows[0].cells[0], table.rows[0].cells[1]
        dot.width = Inches(0.22)
        text_cell.width = Inches(6.28)
        set_cell_margins(dot, top=20, bottom=0, left=40, right=0)
        set_cell_margins(text_cell, top=20, bottom=40, left=0, right=40)
        style_run(dot.paragraphs[0].add_run("●"), size=9, bold=True, color=C["gold"])
        style_run(text_cell.paragraphs[0].add_run(item), size=10, color=C["navy"])


def add_screenshot_placeholder(doc: Document, label: str) -> None:
    # Label chip
    chip = doc.add_table(rows=1, cols=1)
    set_table_width(chip, 6.5)
    remove_table_borders(chip)
    chip_cell = chip.rows[0].cells[0]
    set_cell_shading(chip_cell, C["navy"])
    set_cell_margins(chip_cell, top=60, bottom=60, left=140, right=140)
    cp = chip_cell.paragraphs[0]
    cr = cp.add_run(f"  SCREENSHOT  ·  {label}")
    style_run(cr, size=9, bold=True, color=C["gold_soft"])

    # Placeholder area
    box = doc.add_table(rows=1, cols=1)
    set_table_width(box, 6.5)
    cell = box.rows[0].cells[0]
    set_cell_shading(cell, C["light"])
    set_cell_borders(cell, C["border"], "8")
    set_cell_margins(cell, top=320, bottom=320, left=200, right=200)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    logo_path = LOGO_DIR / "unmul-20260408145731-e033c2.png"
    if logo_path.exists():
        para.add_run().add_picture(str(logo_path), width=Inches(0.55))
        para.add_run("\n")

    r1 = para.add_run("Sisipkan screenshot di sini")
    style_run(r1, size=12, bold=True, color=C["muted"])
    para.add_run("\n")
    r2 = para.add_run("Klik kanan area ini → Sisipkan → Gambar dari perangkat")
    style_run(r2, size=9, italic=True, color=C["muted"])

    doc.add_paragraph().paragraph_format.space_after = Pt(10)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(C["navy"])
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15


def setup_header_footer(section, *, cover: bool = False) -> None:
    section.different_first_page_header_footer = True

    # Header halaman isi
    header = section.header
    header.is_linked_to_previous = False
    htable = header.add_table(rows=1, cols=3, width=Inches(6.5))
    htable.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(htable)
    left, mid, right = htable.rows[0].cells
    left.width, mid.width, right.width = Inches(1.2), Inches(4.1), Inches(1.2)

    unmul = LOGO_DIR / "unmul-20260408145731-e033c2.png"
    if unmul.exists():
        lp = left.paragraphs[0]
        lp.add_run().add_picture(str(unmul), width=Inches(0.42))

    mp = mid.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(mp.add_run("PORTAL ORASI ILMIAH GURU BESAR"), size=8, bold=True, color=C["navy"])
    mp.add_run("\n")
    style_run(mp.add_run("Universitas Mulawarman"), size=7, color=C["muted"])

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(rp.add_run("Dok. Fitur"), size=7, bold=True, color=C["gold"])

    # Gold line under header
    line_p = header.add_paragraph()
    line_p.paragraph_format.space_before = Pt(2)
    line_p.paragraph_format.space_after = Pt(0)
    line_run = line_p.add_run("█" * 90)
    style_run(line_run, size=4, color=C["gold"])

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(
        fp.add_run("Portal Orasi UNMUL  ·  Dokumentasi Fitur  ·  "),
        size=7,
        color=C["muted"],
    )
    style_run(fp.add_run("Kuning · Navy · Emas"), size=7, bold=True, color=C["gold"])

    # Cover: kosongkan header/footer pertama
    if cover:
        first_header = section.first_page_header
        first_header.is_linked_to_previous = False
        if first_header.paragraphs:
            first_header.paragraphs[0].clear()
        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False
        if first_footer.paragraphs:
            first_footer.paragraphs[0].clear()


def add_cover_page(doc: Document) -> None:
    # Top gold bar
    top = doc.add_table(rows=1, cols=1)
    set_table_width(top, 6.8)
    remove_table_borders(top)
    set_cell_shading(top.rows[0].cells[0], C["gold"])
    top.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(5)

    # Navy hero block
    hero = doc.add_table(rows=1, cols=1)
    set_table_width(hero, 6.8)
    remove_table_borders(hero)
    hcell = hero.rows[0].cells[0]
    set_cell_shading(hcell, C["navy_dark"])
    set_cell_margins(hcell, top=280, bottom=200, left=200, right=200)

    hp = hcell.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    unmul = LOGO_DIR / "unmul-20260408145731-e033c2.png"
    if unmul.exists():
        hp.add_run().add_picture(str(unmul), width=Inches(1.15))
    hp.add_run("\n\n")

    t1 = hp.add_run("DOKUMENTASI FITUR")
    style_run(t1, size=26, bold=True, color=C["white"])
    hp.add_run("\n")
    t2 = hp.add_run("PORTAL ORASI ILMIAH GURU BESAR")
    style_run(t2, size=16, bold=True, color=C["gold_soft"])
    hp.add_run("\n")
    t3 = hp.add_run("UNIVERSITAS MULAWARMAN")
    style_run(t3, size=13, bold=True, color=C["white"])
    hp.add_run("\n\n")
    t4 = hp.add_run("Website Publik  ·  Panel Admin  ·  Si Ora Chatbot")
    style_run(t4, size=11, color=C["gold"])

    # Brand logos strip
    logos = doc.add_table(rows=1, cols=6)
    set_table_width(logos, 6.8)
    remove_table_borders(logos)
    for i, (fname, alt) in enumerate(BRAND_LOGOS):
        cell = logos.rows[0].cells[i]
        set_cell_shading(cell, C["white"])
        set_cell_margins(cell, top=80, bottom=80, left=40, right=40)
        path = LOGO_DIR / fname
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            cp.add_run().add_picture(str(path), width=Inches(0.72))
        cp.add_run("\n")
        style_run(cp.add_run(alt.split()[0]), size=6, color=C["muted"])

    # Meta card
    meta = doc.add_table(rows=1, cols=1)
    set_table_width(meta, 6.8)
    remove_table_borders(meta)
    mcell = meta.rows[0].cells[0]
    set_cell_shading(mcell, C["gold_pale"])
    set_cell_margins(mcell, top=160, bottom=160, left=200, right=200)
    set_cell_borders(mcell, C["gold"], "6")
    mp = mcell.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    today = date.today().strftime("%d %B %Y")
    style_run(mp.add_run(f"Versi dokumen: {today}\n"), size=10, bold=True, color=C["navy"])
    style_run(mp.add_run("Repositori: github.com/asyaress/orasi\n"), size=9, color=C["muted"])
    style_run(mp.add_run("Tempatkan screenshot pada setiap kotak abu bertanda SCREENSHOT"), size=9, italic=True, color=C["blue"])

    # Si Ora avatar on cover
    avatar = PUBLIC / "avatar-chat.png"
    if avatar.exists():
        ap = doc.add_paragraph()
        ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ap.add_run().add_picture(str(avatar), width=Inches(0.85))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(cap.add_run("Termasuk panduan fitur Si Ora"), size=9, bold=True, color=C["navy"])

    doc.add_page_break()


def add_toc_page(doc: Document) -> None:
    add_themed_heading(doc, "Daftar Isi", 1)
    add_info_box(
        doc,
        "Dokumen ini menggunakan tema visual Portal Orasi (navy #18213A, emas #EFB12C). "
        "Setiap bagian memiliki kotak SCREENSHOT untuk dokumentasi visual Anda.",
    )
    toc_items = [
        ("01", "Ringkasan sistem"),
        ("02", "Website publik — navigasi & komponen global"),
        ("03", "Halaman Beranda — semua section"),
        ("04", "Halaman Guru Besar — arsip & detail profil"),
        ("05", "Halaman Daftar Orasi"),
        ("06", "Halaman Video Orasi"),
        ("07", "Halaman Dokumen Orasi"),
        ("08", "Halaman Statistik"),
        ("09", "Si Ora — Chatbot"),
        ("10", "Autentikasi admin (Login & 2FA)"),
        ("11", "Panel Admin — semua menu & form input"),
        ("12", "Alur kerja admin (workflow)"),
        ("—", "Lampiran checklist screenshot"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, 6.5)
    hdr = table.rows[0].cells
    for i, title in enumerate(["No", "Bagian", "Halaman"]):
        set_cell_shading(hdr[i], C["navy"])
        set_cell_margins(hdr[i], top=80, bottom=80, left=100, right=100)
        style_run(hdr[i].paragraphs[0].add_run(title), size=9, bold=True, color=C["white"])
    for num, label in toc_items:
        row = table.add_row().cells
        for idx, val in enumerate([num, label, ""]):
            set_cell_margins(row[idx], top=60, bottom=60, left=100, right=100)
            if idx == 0:
                set_cell_shading(row[idx], C["gold_pale"])
                style_run(row[idx].paragraphs[0].add_run(val), size=9, bold=True, color=C["navy"])
            else:
                style_run(row[idx].paragraphs[0].add_run(val), size=9, color=C["navy"])
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    setup_header_footer(section, cover=True)

    add_cover_page(doc)
    add_toc_page(doc)

    # --- 1 Ringkasan ---
    add_themed_heading(doc, "1. Ringkasan Sistem", 1)
    add_paragraph_text(
        doc,
        "Portal Orasi Ilmiah Guru Besar UNMUL menghimpun agenda orasi, profil guru besar, "
        "video YouTube, dokumen akademik, dan statistik dalam satu website publik. "
        "Data dikelola melalui panel admin dengan autentikasi dua faktor (2FA).",
    )
    add_bullets(doc, [
        "Framework: Laravel 12 (PHP 8.2+)",
        "Database: guru besar, orasi ilmiah, fakultas, prodi",
        "Media: storage/app/public (foto, banner, dokumen)",
        "Deploy: php artisan orasi:sync-public-html",
        "Urutan guru besar: TMT terawal di atas (per tahun)",
        "Tanggal publik: Bahasa Indonesia (contoh: 01 Agustus 2020)",
    ])

    # --- 2 Global ---
    add_themed_heading(doc, "2. Website Publik — Navigasi & Komponen Global", 1)

    add_themed_heading(doc, "2.1 Header / Navbar", 2)
    add_bullets(doc, [
        "Logo institusi (6 logo resmi UNMUL & Kemdikbud)",
        "Menu: Beranda · Guru Besar · Daftar Orasi · Video · Dokumen · Statistik",
        "Navbar kuning emas (#EFB12C) di atas halaman",
        "Responsif: menu hamburger di mobile",
    ])
    add_screenshot_placeholder(doc, "Header desktop — navbar kuning + logo lengkap")

    add_themed_heading(doc, "2.2 Footer", 2)
    add_bullets(doc, [
        "Logo institusi + judul portal",
        "Akses cepat & navigasi akademik",
        "Informasi kampus",
    ])
    add_screenshot_placeholder(doc, "Footer — seluruh bagian")

    add_themed_heading(doc, "2.3 Preloader & Animasi", 2)
    add_bullets(doc, ["Garis progress emas-biru tipis", "Fade-in halus saat halaman siap"])
    add_screenshot_placeholder(doc, "Preloader — garis progress atas")

    add_themed_heading(doc, "2.4 Halaman Error 503", 2)
    add_bullets(doc, ["Maintenance mode / server unavailable", "Tema navy-emas + logo UNMUL"])
    add_screenshot_placeholder(doc, "Halaman 503 Service Unavailable")
    doc.add_page_break()

    # --- 3 Home ---
    add_themed_heading(doc, "3. Halaman Beranda (/)", 1)
    add_info_box(doc, "URL: / — Ringkasan 7 section. Tiap section punya tombol «Lihat Selengkapnya».")

    sections_home = [
        ("3.1 Hero (#beranda)", ["Video YouTube lazy", "Judul portal", "Counter total guru besar"], "Hero beranda"),
        ("3.2 Tentang Guru Besar (#tentang-guru-besar)", ["Definisi guru besar (UU)", "Background foto", "Tombol scroll"], "Section tentang guru besar"),
        ("3.3 Guru Besar (#guru-besar)", ["Slider poster 4/slide per tahun", "Klik → detail profil"], "Slider guru besar di home"),
        ("3.4 Agenda (#daftar-orasi)", ["Kartu banner agenda (4 highlight)", "Klik → filter guru besar"], "Section agenda di home"),
        ("3.5 Video (#video-orasi)", ["Slider thumbnail YouTube per tahun"], "Section video di home"),
        ("3.6 Dokumen (#dokumen-orasi)", ["Kartu naskah & presentasi"], "Section dokumen di home"),
        ("3.7 Statistik (#statistik)", ["Chart gender donut", "Chart fakultas bar"], "Section statistik di home"),
    ]
    for title, bullets, ss in sections_home:
        add_themed_heading(doc, title, 2)
        add_bullets(doc, bullets)
        add_screenshot_placeholder(doc, ss)
    add_screenshot_placeholder(doc, "Beranda — full page scroll panjang")
    doc.add_page_break()

    # --- 4 Guru Besar ---
    add_themed_heading(doc, "4. Halaman Guru Besar (/guru-besar)", 1)
    add_bullets(doc, ["Arsip accordion per tahun", "Urutan TMT terawal di atas", "Filter ?orasi={id}"])
    add_screenshot_placeholder(doc, "Guru besar — accordion tahun terbuka")

    add_themed_heading(doc, "4.1 Detail Profil (/guru-besar/{id})", 2)
    for name, desc in [
        ("Hero profil", "Nama, badge, poster foto"),
        ("Data Guru Besar", "TMT, fakultas, prodi, judul orasi"),
        ("Pratinjau Piagam", "Embed PDF/gambar"),
        ("Dokumen Orasi", "Naskah, PPT, piagam, sertifikat"),
        ("Video Orasi", "Embed YouTube"),
        ("Unduh Paket", "ZIP foto + dokumen"),
        ("Guru Besar Lain", "Slider tahun sama"),
    ]:
        add_bullets(doc, [f"{name}: {desc}"])
        add_screenshot_placeholder(doc, f"Detail — {name}")
    doc.add_page_break()

    # --- 5-8 ---
    for title, bullets, ss in [
        ("5. Daftar Orasi (/daftar-orasi)", ["Semua agenda per tahun", "Status published/archived"], "Halaman daftar orasi"),
        ("6. Video Orasi (/video-orasi)", ["Accordion per tahun", "Grid thumbnail YouTube"], "Halaman video orasi"),
        ("7. Dokumen Orasi (/dokumen-orasi)", ["Kartu dokumen per guru", "Unduh gabungan per tahun"], "Halaman dokumen orasi"),
        ("8. Statistik (/statistik)", ["4 chart + ringkasan angka", "Raihan SK 2019–2025"], "Halaman statistik lengkap"),
    ]:
        add_themed_heading(doc, title, 1)
        add_bullets(doc, bullets)
        add_screenshot_placeholder(doc, ss)
    doc.add_page_break()

    # --- 9 Si Ora ---
    add_themed_heading(doc, "9. Si Ora — Asisten Chatbot", 1)
    avatar = PUBLIC / "avatar-chat.png"
    if avatar.exists():
        ap = doc.add_paragraph()
        ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ap.add_run().add_picture(str(avatar), width=Inches(1.0))
    add_info_box(doc, "Si Ora membantu penelusuran agenda, profil guru besar, statistik, video, dan dokumen via chat di beranda.")
    add_themed_heading(doc, "9.1 Tampilan UI", 2)
    add_bullets(doc, ["Tombol floating avatar", "Panel chat + suggestions", "Indikator Mengetik..."])
    for ss in ["Si Ora — tombol floating", "Si Ora — panel terbuka", "Si Ora — contoh jawaban"]:
        add_screenshot_placeholder(doc, ss)
    add_themed_heading(doc, "9.2 Kemampuan", 2)
    add_bullets(doc, [
        "Statistik & jumlah guru besar per fakultas/tahun",
        "Pencarian nama guru besar",
        "Info TMT, bidang ilmu, judul orasi",
        "Agenda, video, dokumen, navigasi portal",
    ])
    doc.add_page_break()

    # --- 10 Auth ---
    add_themed_heading(doc, "10. Autentikasi Admin", 1)
    add_themed_heading(doc, "10.1 Login", 2)
    add_screenshot_placeholder(doc, "Halaman login admin")
    add_themed_heading(doc, "10.2 Two-Factor (2FA)", 2)
    add_screenshot_placeholder(doc, "Setup 2FA — QR code")
    add_screenshot_placeholder(doc, "Challenge 2FA — input 6 digit")
    doc.add_page_break()

    # --- 11 Admin ---
    add_themed_heading(doc, "11. Panel Admin (/admin)", 1)
    add_screenshot_placeholder(doc, "Layout admin — sidebar kuning + konten")

    for title, desc, ss in [
        ("11.1 Dashboard", "Kartu statistik + shortcut", "Dashboard admin"),
        ("11.2 Orasi Ilmiah — Index", "Tabel event per tahun", "Index orasi ilmiah"),
        ("11.3 Orasi Ilmiah — Form", "Field tahun, judul, tanggal, banner", "Form orasi ilmiah"),
        ("11.4 Penugasan Guru Besar", "Drag-and-drop assign", "Panel penugasan orasi"),
        ("11.5 Guru Besar — Index", "Tabel master data", "Index guru besar"),
        ("11.6 Guru Besar — Form", "Identitas, media, foto", "Form guru besar"),
        ("11.7 Statistic", "Grafik admin", "Statistic admin"),
        ("11.8 Pengumuman", "CRUD pengumuman", "Pengumuman admin"),
        ("11.9 Arsip", "Tabel arsip event", "Arsip admin"),
        ("11.10 Security", "Device 2FA", "Security admin"),
        ("11.11 Fakultas", "Master fakultas", "Fakultas admin"),
        ("11.12 Prodi", "Master prodi", "Prodi admin"),
    ]:
        add_themed_heading(doc, title, 2)
        add_paragraph_text(doc, desc, size=10, color=C["muted"])
        add_screenshot_placeholder(doc, ss)

    add_themed_heading(doc, "11.13 Field Input — Orasi Ilmiah", 2)
    add_bullets(doc, ["Tahun*, Judul*, Tanggal*, Jenis (Luring/Daring)", "Pendaftaran mulai/selesai", "Status*, Banner upload"])
    add_screenshot_placeholder(doc, "Form orasi — close-up semua field")

    add_themed_heading(doc, "11.14 Field Input — Guru Besar", 2)
    add_bullets(doc, [
        "Identitas: Pegawai ID, Sumber, Nama*, Gender*, Bidang, Judul, TMT",
        "Fakultas & Prodi (+ snapshot manual)",
        "Media: YouTube, naskah, PPT, piagam, sertifikat",
        "Foto: mode SVG/PNG overlay + upload",
    ])
    add_screenshot_placeholder(doc, "Form guru besar — identitas & media")
    doc.add_page_break()

    # --- 12 Workflow ---
    add_themed_heading(doc, "12. Alur Kerja Admin", 1)
    for i, step in enumerate([
        "Master Fakultas & Prodi",
        "Buat Orasi Ilmiah tahun baru",
        "Upload banner event",
        "Input Guru Besar + media",
        "Penugasan drag-drop di detail orasi",
        "Publish → tampil di portal",
        "Deploy: git pull → migrate → sync-public-html",
    ], 1):
        table = doc.add_table(rows=1, cols=2)
        set_table_width(table, 6.5)
        remove_table_borders(table)
        num, txt = table.rows[0].cells[0], table.rows[0].cells[1]
        num.width, txt.width = Inches(0.55), Inches(5.95)
        set_cell_shading(num, C["gold"])
        set_cell_margins(num, top=80, bottom=80, left=60, right=60)
        style_run(num.paragraphs[0].add_run(str(i)), size=12, bold=True, color=C["navy"])
        set_cell_shading(txt, C["light"])
        set_cell_margins(txt, top=80, bottom=80, left=140, right=140)
        style_run(txt.paragraphs[0].add_run(step), size=10, color=C["navy"])

    add_themed_heading(doc, "Lampiran — Checklist Screenshot", 1)
    checklist = [
        "Header & Footer", "Preloader", "Home 7 section",
        "Guru Besar + Detail", "Daftar/Video/Dokumen/Statistik",
        "Si Ora (3 state)", "Login & 2FA", "Semua halaman admin",
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, 6.5)
    for i, h in enumerate(["Item", "✓ Selesai"]):
        set_cell_shading(table.rows[0].cells[i], C["navy"])
        style_run(table.rows[0].cells[i].paragraphs[0].add_run(h), size=9, bold=True, color=C["white"])
    for idx, item in enumerate(checklist):
        row = table.add_row().cells
        bg = C["white"] if idx % 2 else C["light"]
        for cell in row:
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        style_run(row[0].paragraphs[0].add_run(item), size=9, color=C["navy"])
        style_run(row[1].paragraphs[0].add_run("☐"), size=11, color=C["gold"])

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
